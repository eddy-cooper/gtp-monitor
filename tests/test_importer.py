"""Tests for the workbook importer: the pure parsing helpers, and the
demo workbook driven through the full parse/import path. The demo
workbook deliberately reproduces the awkward habits of a real
hand-maintained operational spreadsheet (varying day-1 rows, formula
strings, free-text timestamps) -- see docs/IMPORTER_NOTES.md.
"""

import json
import sqlite3
import sys
from pathlib import Path

import pytest

from gtp.importer import (
    _parse_mains_readings,
    import_workbook,
    parse_snapshot_time,
    parse_workbook,
    safe_eval_arithmetic,
)

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import make_demo_data as mdd  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
DEMO_WORKBOOK = REPO_ROOT / "demo" / "Demo Operational Record.xlsx"
EXPECTED_FIGURES = json.loads(
    (REPO_ROOT / "demo" / "expected_figures.json").read_text(encoding="utf-8")
)


# --- pure helpers -----------------------------------------------------------


def test_safe_eval_arithmetic_handles_the_workbook_formula_shapes():
    assert safe_eval_arithmetic("=1655.1-1650.0+(0/30*5.3)") == pytest.approx(5.1)
    assert safe_eval_arithmetic("=5.115+2.23") == pytest.approx(7.345)
    assert safe_eval_arithmetic("2+3*4") == pytest.approx(14)
    assert safe_eval_arithmetic("-(1.5)") == pytest.approx(-1.5)


def test_safe_eval_arithmetic_rejects_anything_but_arithmetic():
    with pytest.raises(ValueError):
        safe_eval_arithmetic("=SUM(A1:A3)")
    with pytest.raises(ValueError):
        safe_eval_arithmetic("__import__('os')")
    with pytest.raises(ValueError):
        safe_eval_arithmetic("=G13*E$11")  # cell references are not arithmetic


def test_parse_mains_readings_extracts_end_and_start():
    end, start = _parse_mains_readings("=1655.1-1650.0+(0/30*5.3)")
    assert end == pytest.approx(1655.1)
    assert start == pytest.approx(1650.0)


def test_parse_snapshot_time_clean():
    iso, note = parse_snapshot_time("30 Jun @ 0915am", month=6, year=2026)
    assert iso == "2026-06-30T09:15:00"
    assert note is None


def test_parse_snapshot_time_midnight_and_its_typo():
    iso, note = parse_snapshot_time("31 Oct @ midnight", month=10, year=2025)
    assert iso == "2025-10-31T00:00:00"
    assert note is None
    iso, note = parse_snapshot_time("31 Mar @ midnght", month=3, year=2026)
    assert iso == "2026-03-31T00:00:00"
    assert note is None  # the recurring sheet typo still parses


def test_parse_snapshot_time_pm_and_bare_hour():
    iso, _ = parse_snapshot_time("30 Apr @ 5pm", month=4, year=2026)
    assert iso == "2026-04-30T17:00:00"
    iso, _ = parse_snapshot_time("30 Nov @ 1145pm", month=11, year=2025)
    assert iso == "2025-11-30T23:45:00"


def test_parse_snapshot_time_previous_month_spillover():
    # a sheet's opening snapshot is usually dated in the previous month
    iso, _ = parse_snapshot_time("30 Sep @ midnight", month=10, year=2025)
    assert iso == "2025-09-30T00:00:00"


def test_parse_snapshot_time_failure_keeps_raw_and_flags():
    iso, note = parse_snapshot_time("???", month=3, year=2026)
    assert iso == "2026-03-01T00:00:00"  # midnight day-1 fallback
    assert note is not None and "???" in note


# --- the demo workbook through parse_workbook -------------------------------


@pytest.fixture(scope="module")
def sheets():
    return {s.sheet_name: s for s in parse_workbook(DEMO_WORKBOOK)}


def test_every_sheet_parses_with_the_right_month(sheets):
    assert len(sheets) == EXPECTED_FIGURES["sheets"]
    assert sheets["Oct 25"].year == 2025 and sheets["Oct 25"].month == 10
    assert sheets["Aug 26"].year == 2026 and sheets["Aug 26"].month == 8


def test_day1_row_varies_and_is_found_on_every_sheet(sheets):
    import openpyxl

    wb = openpyxl.load_workbook(DEMO_WORKBOOK)
    found_rows = set()
    for spec in mdd.MONTHS:
        ws = wb[spec.sheet]
        assert ws[f"B{spec.day1_row}"].value == 1
        found_rows.add(spec.day1_row)
    assert len(found_rows) > 3  # genuinely varies, not one fixed offset


def test_daily_rows_cover_every_calendar_day(sheets):
    assert len(sheets["Oct 25"].daily_readings) == 31
    assert len(sheets["Feb 26"].daily_readings) == 28
    # August has all 31 day rows; days after the 14th are placeholders
    august = sheets["Aug 26"].daily_readings
    assert len(august) == 31
    assert august[13].fit0101_reading is not None
    assert august[14].fit0101_reading is None


def test_totals_match_recorded_figures(sheets):
    for key in ("2025-10-01..2025-10-31", "2026-07-01..2026-07-31"):
        expected = EXPECTED_FIGURES["periods"][key]
        sheet = next(
            s for s in sheets.values()
            if f"{s.year:04d}-{s.month:02d}-01" == expected["period_start"]
        )
        filled = [d for d in sheet.daily_readings if d.fit0101_reading is not None]
        total_in = sum(d.fit0101_reading - d.fit0101_error for d in filled)
        total_out = sum(d.fit0501 for d in filled)
        assert total_in == pytest.approx(expected["total_in"], rel=1e-9)
        assert total_out == pytest.approx(expected["total_out"], rel=1e-9)


def test_mains_formula_string_is_evaluated_not_trusted(sheets):
    oct_sheet = sheets["Oct 25"]
    assert oct_sheet.mains_used_reported == pytest.approx(5.1)
    assert oct_sheet.end_snapshot.mains_meter == pytest.approx(1655.1)
    assert oct_sheet.start_snapshot.mains_meter == pytest.approx(1650.0)


def test_plc_freeze_comments_mark_days_estimated(sheets):
    flagged = [
        d.date
        for s in sheets.values()
        for d in s.daily_readings
        if d.is_estimated
    ]
    assert sorted(flagged) == EXPECTED_FIGURES["estimated_dates"]


def test_unparsed_service_cell_becomes_a_warning(sheets):
    warnings = [w for s in sheets.values() for w in s.warnings]
    assert len(warnings) == 1
    assert "OFFLINE" in warnings[0]


def test_config_change_is_read_per_sheet(sheets):
    era_a = EXPECTED_FIGURES["config_eras"]["2025-10-01"]
    era_b = EXPECTED_FIGURES["config_eras"]["2026-04-01"]
    assert sheets["Mar 26"].batch_volume_kl == pytest.approx(era_a["batch_volume_kl"])
    assert sheets["Apr 26"].batch_volume_kl == pytest.approx(era_b["batch_volume_kl"])
    assert sheets["Apr 26"].chem_ferrous_factor == pytest.approx(era_b["ferrous"])


# --- import into the database ----------------------------------------------


def test_import_writes_and_shares_boundary_snapshots(tmp_path):
    db_path = tmp_path / "gtp.db"
    summary = import_workbook(DEMO_WORKBOOK, db_path)
    assert summary.sheets_imported == EXPECTED_FIGURES["sheets"]
    assert summary.days_imported == EXPECTED_FIGURES["days_total"]

    conn = sqlite3.connect(db_path)
    (snapshots,) = conn.execute("SELECT COUNT(*) FROM balance_snapshot").fetchone()
    # 11 sheets share their month boundaries -> 12 distinct snapshots
    assert snapshots == EXPECTED_FIGURES["distinct_snapshots"]


def test_import_is_idempotent(tmp_path):
    db_path = tmp_path / "gtp.db"
    import_workbook(DEMO_WORKBOOK, db_path)
    import_workbook(DEMO_WORKBOOK, db_path)

    conn = sqlite3.connect(db_path)
    for table, expected in (
        ("daily_reading", EXPECTED_FIGURES["days_total"]),
        ("balance_snapshot", EXPECTED_FIGURES["distinct_snapshots"]),
        ("config_history", EXPECTED_FIGURES["sheets"]),
    ):
        (count,) = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()
        assert count == expected, table
